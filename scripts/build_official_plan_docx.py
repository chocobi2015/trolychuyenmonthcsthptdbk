import json
import os
import re
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = 'Times New Roman'

def set_run_font(run, name=FONT_NAME, size_pt=14, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    # Ensure eastAsia/complex script font is also Times New Roman
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:cs="{name}" w:eastAsia="{name}"/>')
    rPr.append(rFonts)

def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        '<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                        '<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                        '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                        '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                        '<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                        '<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                        '</w:tblBorders>')
    tblPr.append(borders)

def format_paragraph_text(p, text, default_bold=False, default_italic=False, size_pt=14):
    parts = re.split(r'(\*\*.*?\*\*)', text)
    is_numbered = bool(re.match(r'^\d+(\.\d+)*\.', text.strip()))
    
    for idx, part in enumerate(parts):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            inner = part[2:-2]
            if inner:
                run = p.add_run(inner)
                set_run_font(run, FONT_NAME, size_pt, bold=True, italic=default_italic)
        else:
            bold = default_bold
            if is_numbered and idx == 0 and ':' in part:
                bold = True
            run = p.add_run(part)
            set_run_font(run, FONT_NAME, size_pt, bold=bold, italic=default_italic)

def build_plan_document(output_path):
    json_path = os.path.join(os.getcwd(), 'src/data/persistedDocuments.json')
    with open(json_path, 'r', encoding='utf-8') as f:
        docs = json.load(f)
    
    doc_data = next((d for d in docs if d['id'] == 'doc-kh-ktgk1-2026'), None)
    if not doc_data:
        raise ValueError("doc-kh-ktgk1-2026 not found in persistedDocuments.json")

    doc = docx.Document()
    
    # 1. Page Margins & Size: A4, Top 2.0cm, Bottom 2.0cm, Left 2.5cm, Right 1.8cm
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.8)

    # 2. Header Table: 2 columns, borderless
    # Printable width: 21.0 - 2.5 - 1.8 = 16.7 cm
    # Left col: 7.2 cm, Right col: 9.5 cm
    header_table = doc.add_table(rows=2, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    remove_table_borders(header_table)

    col_widths = [Cm(7.2), Cm(9.5)]
    for row in header_table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    # Row 0 Left: Cơ quan ban hành
    c00 = header_table.cell(0, 0)
    p00_1 = c00.paragraphs[0]
    p00_1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p00_1.paragraph_format.space_after = Pt(2)
    r = p00_1.add_run('SỞ GDĐT TỈNH ĐỒNG THÁP')
    set_run_font(r, FONT_NAME, 12, bold=False)

    p00_2 = c00.add_paragraph()
    p00_2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p00_2.paragraph_format.space_after = Pt(2)
    r = p00_2.add_run('TRƯỜNG THCS VÀ THPT\nĐỐC BINH KIỀU')
    set_run_font(r, FONT_NAME, 12, bold=True)

    p00_3 = c00.add_paragraph()
    p00_3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p00_3.paragraph_format.space_after = Pt(4)
    r = p00_3.add_run('───────')
    set_run_font(r, FONT_NAME, 7, bold=False)

    # Row 0 Right: Quốc hiệu & Tiêu ngữ
    c01 = header_table.cell(0, 1)
    p01_1 = c01.paragraphs[0]
    p01_1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p01_1.paragraph_format.space_after = Pt(2)
    r = p01_1.add_run('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM')
    set_run_font(r, FONT_NAME, 12, bold=True)

    p01_2 = c01.add_paragraph()
    p01_2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p01_2.paragraph_format.space_after = Pt(2)
    r = p01_2.add_run('Độc lập - Tự do - Hạnh phúc')
    set_run_font(r, FONT_NAME, 13, bold=True)

    p01_3 = c01.add_paragraph()
    p01_3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p01_3.paragraph_format.space_after = Pt(4)
    r = p01_3.add_run('────────────────')
    set_run_font(r, FONT_NAME, 7, bold=False)

    # Row 1 Left: Số ký hiệu
    c10 = header_table.cell(1, 0)
    p10 = c10.paragraphs[0]
    p10.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p10.paragraph_format.space_before = Pt(2)
    p10.paragraph_format.space_after = Pt(4)
    r = p10.add_run(doc_data.get('documentNumber', 'Số: 45/KH-THCS&THPTĐBK'))
    set_run_font(r, FONT_NAME, 12, bold=False)

    # Row 1 Right: Địa danh & Ngày tháng
    c11 = header_table.cell(1, 1)
    p11 = c11.paragraphs[0]
    p11.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p11.paragraph_format.space_before = Pt(2)
    p11.paragraph_format.space_after = Pt(4)
    r = p11.add_run(doc_data.get('signDate', 'Tháp Mười, ngày 16 tháng 10 năm 2026'))
    set_run_font(r, FONT_NAME, 13, bold=False, italic=True)

    # 3. Spacing before Title
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(6)
    sp.paragraph_format.space_after = Pt(2)

    # Title: KẾ HOẠCH
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run(doc_data.get('title', 'KẾ HOẠCH').upper())
    set_run_font(r_title, FONT_NAME, 15, bold=True)

    # Subtitle
    if doc_data.get('subTitle'):
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(2)
        p_sub.paragraph_format.space_after = Pt(2)
        r_sub = p_sub.add_run(doc_data['subTitle'])
        set_run_font(r_sub, FONT_NAME, 14, bold=True)

        p_line = doc.add_paragraph()
        p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_line.paragraph_format.space_before = Pt(0)
        p_line.paragraph_format.space_after = Pt(8)
        r_line = p_line.add_run('───────────────')
        set_run_font(r_line, FONT_NAME, 7, bold=False)

    # 4. Legal Bases (Căn cứ...)
    for lb in doc_data.get('legalBases', []):
        text = lb if lb.startswith('Căn cứ') else f'Căn cứ {lb}'
        p_lb = doc.add_paragraph()
        p_lb.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_lb.paragraph_format.first_line_indent = Cm(1.0)
        p_lb.paragraph_format.line_spacing = 1.15
        p_lb.paragraph_format.space_after = Pt(3)
        r_lb = p_lb.add_run(text)
        set_run_font(r_lb, FONT_NAME, 14, italic=True)

    # 5. Transition sentence
    p_tr = doc.add_paragraph()
    p_tr.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_tr.paragraph_format.first_line_indent = Cm(1.0)
    p_tr.paragraph_format.line_spacing = 1.15
    p_tr.paragraph_format.space_before = Pt(4)
    p_tr.paragraph_format.space_after = Pt(6)
    r_tr = p_tr.add_run('Nay Trường THCS và THPT Đốc Binh Kiều xây dựng Kế hoạch tổ chức kiểm tra, đánh giá giữa học kỳ I năm học 2026 - 2027 với các nội dung chi tiết như sau:')
    set_run_font(r_tr, FONT_NAME, 14)

    # Helper to render sections
    def render_section_content(sec):
        is_appendix = sec['heading'].upper().startswith('PHỤ LỤC')
        
        # Heading
        p_h = doc.add_paragraph()
        p_h.paragraph_format.line_spacing = 1.15
        if is_appendix:
            p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_h.paragraph_format.space_before = Pt(14)
            p_h.paragraph_format.space_after = Pt(8)
            r = p_h.add_run(sec['heading'].upper())
            set_run_font(r, FONT_NAME, 14, bold=True)
        else:
            p_h.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_h.paragraph_format.first_line_indent = Cm(1.0)
            p_h.paragraph_format.space_before = Pt(8)
            p_h.paragraph_format.space_after = Pt(4)
            r = p_h.add_run(sec['heading'])
            set_run_font(r, FONT_NAME, 14, bold=True)

        lines = sec['content'].split('\n')
        table_lines = []

        def flush_table(t_lines):
            if not t_lines:
                return
            parsed = []
            for tl in t_lines:
                if '---' in tl:
                    continue
                cells = [c.strip() for c in tl.split('|')[1:-1]]
                if cells:
                    parsed.append(cells)
            if not parsed:
                return
            
            num_rows = len(parsed)
            num_cols = len(parsed[0])
            tbl = doc.add_table(rows=num_rows, cols=num_cols, style='Table Grid')
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = True

            for r_idx, row_cells in enumerate(parsed):
                row = tbl.rows[r_idx]
                is_header = (r_idx == 0)
                for c_idx, cell_value in enumerate(row_cells):
                    if c_idx < len(row.cells):
                        cell = row.cells[c_idx]
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.15
                        
                        clean_text = cell_value.replace('**', '').strip()
                        is_center = (c_idx == 0 or len(clean_text) <= 6 or (clean_text and clean_text[0].isdigit() and len(clean_text) < 15) or is_header)
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_center else WD_ALIGN_PARAGRAPH.LEFT
                        
                        r = p.add_run(clean_text or '—')
                        set_run_font(r, FONT_NAME, 12, bold=(is_header or '**' in cell_value))

            sp_after = doc.add_paragraph()
            sp_after.paragraph_format.space_before = Pt(4)
            sp_after.paragraph_format.space_after = Pt(4)

        for line in lines:
            trimmed = line.strip()
            if trimmed.startswith('|') and trimmed.endswith('|'):
                table_lines.append(trimmed)
            else:
                if table_lines:
                    flush_table(table_lines)
                    table_lines = []
                if trimmed:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p.paragraph_format.first_line_indent = Cm(1.0)
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(3)
                    format_paragraph_text(p, trimmed, size_pt=14)

        if table_lines:
            flush_table(table_lines)

    # 6. Separate Regular vs Appendix Sections
    regular_secs = [s for s in doc_data.get('sections', []) if not s['heading'].upper().startswith('PHỤ LỤC')]
    appendix_secs = [s for s in doc_data.get('sections', []) if s['heading'].upper().startswith('PHỤ LỤC')]

    for s in regular_secs:
        render_section_content(s)

    # 7. Signature & Recipients Block (footer_table)
    sp_f = doc.add_paragraph()
    sp_f.paragraph_format.space_before = Pt(8)
    sp_f.paragraph_format.space_after = Pt(4)

    footer_table = doc.add_table(rows=1, cols=2)
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    footer_table.autofit = False
    remove_table_borders(footer_table)

    footer_col_widths = [Cm(8.0), Cm(8.7)]
    for r in footer_table.rows:
        for i, w in enumerate(footer_col_widths):
            r.cells[i].width = w

    # Left cell: Nơi nhận
    cf_left = footer_table.cell(0, 0)
    p_nn = cf_left.paragraphs[0]
    p_nn.paragraph_format.space_after = Pt(2)
    r_nn = p_nn.add_run('Nơi nhận:')
    set_run_font(r_nn, FONT_NAME, 12, bold=True, italic=True)

    recipients = doc_data.get('recipients', [
        '- Sở GDĐT Đồng Tháp (báo cáo);',
        '- Ban Giám hiệu (chỉ đạo, thực hiện);',
        '- Các Tổ CM, Tổ VP (thực hiện);',
        '- 03 Điểm trường (thực hiện);',
        '- Ban ĐD Cha mẹ học sinh (phối hợp);',
        '- Đoàn trường, Đội TNTP (phối hợp);',
        '- Website trường (công khai);',
        '- Lưu: VT, CM.'
    ])

    for rc in recipients:
        p_rc = cf_left.add_paragraph()
        p_rc.paragraph_format.space_after = Pt(1)
        p_rc.paragraph_format.line_spacing = 1.15
        text = rc if rc.startswith('-') else f'- {rc}'
        r_rc = p_rc.add_run(text)
        set_run_font(r_rc, FONT_NAME, 11)

    # Right cell: Signer Role & Name
    cf_right = footer_table.cell(0, 1)
    signer_lines = doc_data.get('signerRole', 'KT. HIỆU TRƯỞNG\nPHÓ HIỆU TRƯỞNG').split('\n')
    for idx, sl in enumerate(signer_lines):
        p_sl = cf_right.paragraphs[0] if idx == 0 else cf_right.add_paragraph()
        p_sl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sl.paragraph_format.space_after = Pt(2)
        r_sl = p_sl.add_run(sl.upper())
        size = 13 if 'KT.' in sl else 14
        set_run_font(r_sl, FONT_NAME, size, bold=True)

    # Space for handwritten signature & stamp
    for _ in range(5):
        p_blank = cf_right.add_paragraph()
        p_blank.paragraph_format.space_after = Pt(2)
        r_blank = p_blank.add_run('')
        set_run_font(r_blank, FONT_NAME, 14)

    p_name = cf_right.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_after = Pt(4)
    r_name = p_name.add_run(doc_data.get('signerName', 'Nguyễn Minh Trí'))
    set_run_font(r_name, FONT_NAME, 14, bold=True)

    # 8. Render Appendix Sections (after signature block)
    for s in appendix_secs:
        render_section_content(s)

    # 9. Save document
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f'[SUCCESS] Official Plan DOCX built successfully at {output_path} (size: {os.path.getsize(output_path)} bytes)')

if __name__ == '__main__':
    targets = [
        os.path.join(os.getcwd(), 'public/downloads/KH_KIEM_TRA_GIUA_HK1_2026_2027_THCS_THPT_DOC_BINH_KIEU.docx'),
        os.path.join(os.getcwd(), 'dist/downloads/KH_KIEM_TRA_GIUA_HK1_2026_2027_THCS_THPT_DOC_BINH_KIEU.docx'),
        os.path.join(os.getcwd(), 'KIEM TRA DANH GIA/KH_KIEM_TRA_GIUA_HK1_2026_2027_THCS_THPT_DOC_BINH_KIEU.docx')
    ]
    for target in targets:
        build_plan_document(target)
